#!/usr/bin/env python3
"""Build everything the Information Systems submission portal asks for, from an elsarticle main.tex.

Writes into OUT:
  00_Abstract/Abstract.docx   (required upload in the portal; .txt if python-docx is missing)
  PORTAL_ANSWERS.md           page-by-page answers for submit.elsevier.com/INFOSYS
and prints warnings for anything that will block or contradict a portal answer.

Usage:
  python3 build_portal_kit.py --tex main.tex --out upload_ready [--highlights highlights.txt]
"""
import argparse
import os
import re
import sys

ap = argparse.ArgumentParser()
ap.add_argument("--tex", required=True)
ap.add_argument("--out", required=True)
ap.add_argument("--highlights")
args = ap.parse_args()

tex = open(args.tex, encoding="utf-8").read()
warn = []


def braced(s, i):
    """Return (content, end) of the balanced {...} group starting at s[i] == '{'."""
    assert s[i] == "{", s[i:i + 20]
    depth, j = 0, i
    while j < len(s):
        c = s[j]
        if c == "\\":
            j += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    raise ValueError("unbalanced braces near: " + s[i:i + 60])


def cmd(name, s=tex, start=0):
    m = re.compile(r"\\" + name + r"(?:\[[^\]]*\])?\s*\{").search(s, start)
    if not m:
        return None, None
    return braced(s, m.end() - 1)


bib = re.findall(r"\\bibitem(?:\[[^\]]*\])?\{([^}]+)\}", tex)


def plain(s):
    s = re.sub(r"(?<!\\)%.*", "", s)
    s = re.sub(r"\\cite[pt]?\{([^}]+)\}", lambda m: "[" + ", ".join(
        str(bib.index(k.strip()) + 1) if k.strip() in bib else "?" for k in m.group(1).split(",")) + "]", s)
    s = re.sub(r"\\(?:ref|eqref)\{[^}]*\}", "?", s)
    for _ in range(3):
        s = re.sub(r"\\(?:emph|textbf|textit|texttt|textsc|mbox|url)\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\corref\{[^}]*\}", "", s)
    s = s.replace("---", "\u2014").replace("--", "\u2013").replace("~", " ")
    s = s.replace("``", '"').replace("''", '"').replace("\\&", "&").replace("\\%", "%").replace("\\,", " ")
    s = s.replace("\\_", "_").replace("\\#", "#")
    s = re.sub(r"\$([^$]*)\$", r"\1", s)
    left = re.findall(r"\\[A-Za-z]+", s)
    if left:
        warn.append(f"unconverted LaTeX {sorted(set(left))} in: {s[:80]}")
    return re.sub(r"\s+", " ", s).strip()


def ascii_portal(s):
    """Portal text boxes: keep it plain ASCII so nothing is garbled."""
    return (s.replace("\u2013", "-").replace("\u2014", "-").replace("\u2019", "'")
            .replace("\u201c", '"').replace("\u201d", '"'))


def section(title_regex):
    m = re.search(r"\\section\*?\{(" + title_regex + r")\}", tex, re.I)
    if not m:
        return None, None
    nxt = re.compile(r"\\section\*?\{|\\appendix|\\begin\{thebibliography\}|\\bibliography\{|\\end\{document\}").search(tex, m.end())
    return m.group(1), tex[m.end():nxt.start()]


def paragraphs(block):
    out = []
    for para in re.split(r"\n\s*\n", block.strip()):
        if "\\begin{itemize}" in para:
            pre, rest = para.split("\\begin{itemize}", 1)
            items, post = rest.split("\\end{itemize}", 1)
            if pre.strip():
                out.append(plain(pre))
            out.append("\n".join("- " + plain(it) for it in items.split("\\item")[1:]))
            if post.strip():
                out.append(plain(post))
        else:
            out.append(plain(para))
    return "\n\n".join(p for p in out if p)


# ---------- front matter ----------
title = plain(cmd("title")[0])
abstract = ascii_portal(plain(tex[tex.index("\\begin{abstract}") + 16:tex.index("\\end{abstract}")]))
words = len(abstract.split())
if words > 250:
    warn.append(f"abstract has {words} words (limit 250)")
kw_block = tex[tex.index("\\begin{keyword}") + 15:tex.index("\\end{keyword}")] if "\\begin{keyword}" in tex else ""
keywords = [plain(k) for k in re.split(r"\\sep", kw_block) if plain(k)]
if not 1 <= len(keywords) <= 7:
    warn.append(f"{len(keywords)} keywords (need 1-7)")

affil = {}
for m in re.finditer(r"\\affiliation\[([^\]]+)\]\s*\{", tex):
    body, _ = braced(tex, m.end() - 1)
    fields = {}
    for fm in re.finditer(r"(\w+)\s*=\s*\{", body):
        fields[fm.group(1)], _ = braced(body, fm.end() - 1)
    affil[m.group(1)] = fields

authors = []
for m in re.finditer(r"\\author\[([^\]]*)\]\s*\{", tex):
    name_raw, end = braced(tex, m.end() - 1)
    nxt = tex.find("\\author", end)
    em = re.search(r"\\ead\{([^}]+)\}", tex[end:nxt if nxt > 0 else len(tex)])
    name = plain(name_raw)
    parts = name.split()
    authors.append({
        "first": " ".join(parts[:-1]), "last": parts[-1], "name": name,
        "keys": [k.strip() for k in m.group(1).split(",") if k.strip()],
        "email": em.group(1) if em else "",
        "corresponding": "\\corref" in name_raw,
    })
if not authors:
    warn.append("no \\author[...]{...} entries found")
if sum(a["corresponding"] for a in authors) != 1:
    warn.append("exactly one author should carry \\corref")

credit = {}
_, cr = section(r"CRediT[^}]*")
if cr:
    for m in re.finditer(r"\\textbf\{([^}:]+):?\}:?\s*(.*?)(?=\\textbf\{|$)", cr, re.S):
        roles = plain(m.group(2)).rstrip(".")
        roles = roles.replace("Writing \u2013 review and editing", "Writing \u2013 review & editing")
        credit[plain(m.group(1)).rstrip(":")] = [r.strip() for r in roles.split(",") if r.strip()]
else:
    warn.append("no CRediT section")

fund_title, fund = section(r"Funding|Acknowledg[^}]*")
ci_title, ci = section(r"Declaration of competing interests?")
data_title, data = section(r"Data (?:and code )?availability[^}]*|Data and code availability")
ai_title, ai = section(r"Declaration of generative AI[^}]*")
for label, val in [("Funding", fund), ("Declaration of competing interest", ci), ("Data availability", data), ("Generative AI declaration", ai)]:
    if val is None:
        warn.append(f"section missing: {label}")

ci_text = paragraphs(ci) if ci else ""
no_conflict = bool(re.search(r"no known competing|nothing to declare|declare no competing|no competing (financial )?interests", ci_text, re.I)) and not re.search(r"employ|consult|advis|stock|share|patent|honorar|director|board", ci_text, re.I)

fund_text = paragraphs(fund) if fund else ""
no_funding = bool(re.search(r"did not receive any specific grant", fund_text, re.I))
funders = re.findall(r"supported by ([A-Z][^,.;(]+)", fund_text)
grants = re.findall(r"\[grant numbers? ([^\]]+)\]|grant (?:No\.|number)\s*([A-Za-z0-9-]+)", fund_text)

# dataset deposit for the research-data page
ds = None
for m in re.finditer(r"\\bibitem(?:\[[^\]]*\])?\{([^}]+)\}(.*?)(?=\\bibitem|\\end\{thebibliography\})", tex, re.S):
    body = m.group(2)
    if "[dataset]" in body or "zenodo" in body.lower() or "mendeley" in body.lower() or "figshare" in body.lower():
        url = re.search(r"https://doi\.org/[^\s}]+", body)
        t = re.search(r"``(.+?)''|\u201c(.+?)\u201d", body)
        repo = "Zenodo" if "zenodo" in body.lower() else "Mendeley Data" if "mendeley" in body.lower() else "figshare" if "figshare" in body.lower() else "(choose)"
        ds = {"key": m.group(1), "url": url.group(0).rstrip(".") if url else None, "repo": repo,
              "title": ("Reproducibility materials for \"" + (t.group(1) or t.group(2)) + "\"") if t and "Reproducibility" in body else plain(body)[:200]}
        if ds["url"] and re.search(r"X{4,}|\\zenodoid", ds["url"] + body):
            warn.append("dataset DOI is still a placeholder")
        break
if re.search(r"X{5,}", tex):
    warn.append("placeholder XXXXX found in main.tex")
if data and not ds:
    warn.append("data statement exists but no [dataset]/repository reference was found; Option C needs a deposit with a DOI or a stated reason")

hl = []
if args.highlights:
    hl = [l.strip().lstrip("-\u2022* ").strip() for l in open(args.highlights, encoding="utf-8") if l.strip()]
    if not 3 <= len(hl) <= 5:
        warn.append(f"{len(hl)} highlights (need 3-5)")
    for h in hl:
        if len(h) > 85:
            warn.append(f"highlight over 85 characters ({len(h)}): {h}")

# ---------- outputs ----------
os.makedirs(os.path.join(args.out, "00_Abstract"), exist_ok=True)
try:
    import docx
    from docx.shared import Pt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    d.add_heading(title, level=1)
    d.add_heading("Abstract", level=2)
    d.add_paragraph(abstract)
    cp = d.core_properties
    cp.author = cp.last_modified_by = authors[0]["name"] if authors else ""
    cp.title = "Abstract"
    cp.comments = ""
    abs_path = os.path.join(args.out, "00_Abstract", "Abstract.docx")
    d.save(abs_path)
except ImportError:
    abs_path = os.path.join(args.out, "00_Abstract", "Abstract.txt")
    open(abs_path, "w", encoding="utf-8").write(title + "\n\nAbstract\n\n" + abstract + "\n")
    warn.append("python-docx missing: wrote Abstract.txt (pip install python-docx for .docx)")


def affil_line(k):
    f = affil.get(k, {})
    return ", ".join(x for x in [f.get("organization"), f.get("addressline"), f.get("city"),
                                 f.get("postcode"), f.get("state"), f.get("country")] if x)


def portal_institution(a):
    f = affil.get(a["keys"][0], {}) if a["keys"] else {}
    org = f.get("organization", "")
    # the portal's institution search knows universities, not their schools
    m = re.search(r"([A-Z][\w .&'-]*University[\w .'-]*)$", org) or re.search(r"(University of [\w .'-]+)", org)
    return (m.group(1).strip() if m else org)


L = []
L.append(f"# Portal answers: {title}\n")
L.append("Generated by build_portal_kit.py from main.tex. Paste plain text exactly as shown. The author clicks every Submit/Confirm button.\n")
L.append("## 1. Choose your article type\n**Research paper** (unless the editors invited a survey, reproducibility or retrospective article).\n")
L.append("## 2. Upload your submission files\n")
L.append("Required (upper box):\n")
L.append("| File type | File |\n|---|---|")
L.append(f"| Abstract | `{os.path.relpath(abs_path, args.out)}` |")
L.append("| Manuscript | `01_Manuscript/Manuscript.pdf` (compiled on Overleaf; run verify_final_pdf.py first) |")
if no_conflict:
    L.append("| Declaration of competing interests | none to declare: tick the confirmation box, **or** upload the Declarations-tool file |")
else:
    L.append("| Declaration of competing interests | `06_Declaration_of_interest/declarationStatement.docx`. **Do NOT tick** \"I confirm that no authors ... have any competing ... interests\": the manuscript declares interests. The box greys out once the file is uploaded. |")
L.append("\nOptional (lower box):\n")
L.append("| File type | File |\n|---|---|")
L.append("| LaTeX source files | `02_LaTeX_source/LaTeX_source.zip` |")
figs = sorted(set(re.findall(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}", tex)))
for f in figs:
    L.append(f"| Figure | `03_Figures/{os.path.basename(f)}{'' if os.path.splitext(f)[1] else '.pdf'}` |")
L.append("| Highlights | `04_Highlights/Highlights.docx` (file name must contain \"highlights\") |")
L.append("| Cover letter | `05_Cover_letter/Cover_Letter.docx` |")
L.append("| Supplementary material | optional when everything is in a DOI deposit cited in the paper; keep the portal set all-or-nothing; zips with ~1,700 files fail (\"Something went wrong\") |")
L.append("\nDo not use: Table (tables are in the manuscript), Graphical abstract (not required), Editor annotated manuscript and Response to reviewers (revisions only), Co-submission to Data in Brief/MethodsX (only if one was prepared).\n")
L.append("## 3. Check or add your title and abstract\nThe portal does not extract them from the PDF. Paste:\n")
L.append("**Your title**\n```\n" + ascii_portal(title) + "\n```")
L.append(f"**Your abstract** ({words} words; must end with \"{' '.join(abstract.split()[-3:])}\")\n```\n" + abstract + "\n```\n")
L.append("## 4. Check or add your author details\nOne block per author, in manuscript order; click **Confirm** after each, **Add another author** for the next.\n")
for i, a in enumerate(authors, 1):
    roles = credit.get(a["name"], [])
    if not roles:
        warn.append(f"no CRediT roles found for {a['name']}")
    L.append(f"**Author {i}**\n")
    L.append("| Field | Value |\n|---|---|")
    L.append(f"| First name | `{a['first']}` |\n| Last name | `{a['last']}` |")
    L.append(f"| Institution | search and pick **{portal_institution(a)}** (one per author; the full affiliations stay in the manuscript) |")
    L.append(f"| Email address | `{a['email']}` |")
    L.append(f"| Contributor roles | {', '.join(roles)} ({len(roles)}) |")
    L.append(f"| This is the corresponding author | {'**tick**' if a['corresponding'] else 'leave unticked'} |")
    L.append("\nManuscript affiliations: " + "; ".join(affil_line(k) for k in a["keys"]) + "\n")
L.append("## 5. Add your funder details\n")
if no_funding:
    L.append("**No** (the manuscript states no specific grant).\n")
else:
    L.append("**Yes**: the manuscript has a funding statement, so answering No would contradict it.\n")
    L.append("| Field | Value |\n|---|---|")
    L.append(f"| Funder | {', '.join(f.strip() for f in funders) or '(see statement)'} |")
    g = [x[0] or x[1] for x in grants]
    L.append(f"| Grant number | {', '.join(g) if g else 'leave empty (institutional resources, no grant number); never invent one'} |")
    L.append("\nStatement in the manuscript:\n\n> " + fund_text.replace("\n", "\n> ") + "\n")
L.append("## 6. Decide if you want to publish open access\nAuthor's decision. Subscription: no fee. Gold open access: APC USD 3150 excl. tax (2026), CC BY / CC BY-NC / CC BY-NC-ND. Does not affect review.\n")
L.append("## 7. Oppose reviewers\nOptional. Name only people with a real conflict; leave empty otherwise.\n")
L.append("## 8. Select your classifications\nPick 3-5 from the fixed list (`references/classifications.md`). There is no \"Data Quality\" or \"Data Cleaning\" entry. Choose entries whose reviewers publish in the paper's community; avoid the journal-wide \"Information Systems\" and algorithm-centred entries when the paper proposes no algorithm. Data-management default: Benchmarking, Database Integrity, Data Integration, Database Systems, Relational Databases (swap in Data Warehouse, Data Mining, Query Processing, etc. to fit the topic).\n")
L.append("## 9. Provide additional information\n")
if ds:
    L.append("| Field | Value |\n|---|---|")
    L.append("| Do you want to share data? | **Yes** |")
    L.append(f"| Link to data repository | `{ds['url']}` (printed verbatim in the article; use the DOI, not GitHub) |")
    L.append(f"| Repository name | **{ds['repo']}** |")
    L.append("| Source of data | **Original data** if the deposit holds the authors' own protocols, code and results; Reference data only for third-party datasets |")
    L.append(f"| Title of the data set | `{ascii_portal(ds['title'])}` ({len(ds['title'])} of 200 characters) |")
else:
    L.append("Research data: no deposit found in the bibliography; answer according to the data statement.")
L.append("\nFree Preprint Service (SSRN): author's decision. YES posts the manuscript publicly after desk review and it stays public even if rejected; NO keeps it private (a preprint can still be posted later). Needs all co-authors' consent.\n")
L.append("## 10. Review and submit\nThe author reviews and clicks Submit. Afterwards: confirmation e-mail to the corresponding author; sign in or set the journal-account password from that e-mail to track the manuscript; record the manuscript number.\n")
L.append("## Text the portal may ask for separately\n")
if keywords:
    L.append("**Keywords**: " + "; ".join(keywords) + "\n")
if ci:
    L.append("**Declaration of competing interest**\n\n" + ci_text + "\n")
if cr:
    L.append("**CRediT**\n\n" + "\n".join(f"{k}: {', '.join(v)}." for k, v in credit.items()) + "\n")
if data:
    L.append("**Data availability**\n\n" + paragraphs(data) + "\n")
if ai:
    L.append("**Generative AI declaration**\n\n" + paragraphs(ai) + "\n")
if hl:
    L.append("**Highlights**\n\n" + "\n".join(f"- {h} ({len(h)})" for h in hl) + "\n")

open(os.path.join(args.out, "PORTAL_ANSWERS.md"), "w", encoding="utf-8").write("\n".join(L))
print(f"wrote {abs_path} and {os.path.join(args.out, 'PORTAL_ANSWERS.md')}")
print(f"title ok | abstract {words} words | {len(keywords)} keywords | {len(authors)} authors | "
      f"competing interests: {'none' if no_conflict else 'DECLARED (do not tick the none box)'} | "
      f"funding: {'none' if no_funding else 'yes'} | dataset: {ds['url'] if ds else 'none'}")
for w in warn:
    print("WARN", w)
sys.exit(1 if any(w.startswith(("placeholder", "dataset DOI")) for w in warn) else 0)
