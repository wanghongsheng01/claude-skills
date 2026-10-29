#!/usr/bin/env python3
"""Turn a small Markdown-like text file into an editable .docx for Editorial Manager uploads
(title page, highlights, cover letter, response letter).

  python3 make_docx.py IN.md OUT.docx

Supported: '# ' heading, '## ' subheading, '- ' bullets, blank-line paragraphs, **bold** spans,
^a,b,* superscripts.
Uses python-docx (pip install python-docx); falls back to macOS `textutil` via HTML.
"""
import html
import re
import subprocess
import sys
import pathlib


def blocks(text):
    out, para = [], []
    for line in text.splitlines():
        if not line.strip():
            if para:
                out.append(("p", " ".join(para)))
                para = []
            continue
        if line.startswith("# "):
            out.append(("h1", line[2:].strip()))
        elif line.startswith("## "):
            out.append(("h2", line[3:].strip()))
        elif re.match(r"\s*[-•]\s+", line):
            if para:
                out.append(("p", " ".join(para)))
                para = []
            out.append(("li", re.sub(r"^\s*[-•]\s+", "", line).strip()))
        else:
            para.append(line.strip())
    if para:
        out.append(("p", " ".join(para)))
    return out


def runs(par, s):
    """Add text with **bold** spans and ^a,b,* superscripts (affiliation markers)."""
    for i, part in enumerate(re.split(r"\*\*", s)):
        for j, piece in enumerate(re.split(r"\^([a-z0-9*]+(?:,[a-z0-9*]+)*)", part)):
            if not piece:
                continue
            r = par.add_run(piece)
            r.bold = bool(i % 2)
            r.font.superscript = bool(j % 2)


def with_python_docx(bl, out):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    for kind, s in bl:
        if kind == "h1":
            d.add_heading(s, level=1)
        elif kind == "h2":
            d.add_heading(s, level=2)
        elif kind == "li":
            runs(d.add_paragraph(style="List Bullet"), s)
        else:
            runs(d.add_paragraph(), s)
    d.save(out)


def with_textutil(bl, out):
    parts = ["<html><body style='font-family:Times New Roman;font-size:12pt'>"]
    for kind, s in bl:
        s = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", html.escape(s))
        s = re.sub(r"\^([a-z0-9*]+(?:,[a-z0-9*]+)*)", r"<sup>\1</sup>", s)
        parts.append({"h1": f"<h1>{s}</h1>", "h2": f"<h2>{s}</h2>", "li": f"<ul><li>{s}</li></ul>"}.get(kind, f"<p>{s}</p>"))
    parts.append("</body></html>")
    tmp = pathlib.Path(out).with_suffix(".tmp.html")
    tmp.write_text("\n".join(parts))
    subprocess.run(["textutil", "-convert", "docx", str(tmp), "-output", out], check=True)
    tmp.unlink()


def main():
    src, out = sys.argv[1], sys.argv[2]
    bl = blocks(pathlib.Path(src).read_text())
    try:
        with_python_docx(bl, out)
    except ImportError:
        with_textutil(bl, out)
    print("wrote", out)


if __name__ == "__main__":
    main()
